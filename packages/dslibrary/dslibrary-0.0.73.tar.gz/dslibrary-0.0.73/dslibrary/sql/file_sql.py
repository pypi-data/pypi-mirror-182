import pandas
import dask
import numpy
import re
import io
import openpyxl
import dslibrary
from dslibrary.utils.filechunker import slicer

from .df_iterator import DFIterator
from .exceptions import SqlException, SqlException_SqlProblem
from .misc import table_name_from_table_spec, format_options_for_table
from .modify_sql import embed_parameters
from .parser import find_referenced_tables
from .s2p_errs import translate_pandas_dask_exceptions
from .sql_to_pandas import SqlToPandas
from .streaming_query import streaming_query, ChunkedRowIterator


def file_sql_read(
        open_table_stream: callable, sql: str, parameters: list=None,
        max_rows: int=0, memory_threshold__data_size=None, memory_threshold__max_rows=None,
        mode: str="auto", format_options=None):
    """
    Run SQL against file-like objects containing columnar data in CSV and other formats.

    :param open_table_stream: Function that will open a file-like stream for table data.
    :param sql:          SQL to run.
    :param parameters:   To fill into sql where '%s' occurs.
    :param max_rows:     Maximum # rows to read from files.
    :param memory_threshold__data_size: If data to process is below this size, process in memory.  Measures sum of file sizes.
    :param memory_threshold__max_rows:  If max_rows is below this size, process in memory.
    :param mode:         'pandas', 'stream', 'dask' or 'auto'.
                            'pandas' = load into pandas.DataFrame, e.g. with read_csv().
                            'stream' = do all analysis 'on the fly', using little memory
                            'dask' = use dask instead of pandas - NOTE: dask does not restrain its memory use in many cases
                            'auto' = decide automatically
    :param format_options: Determines file format for a given table.  A mapping from table name to DataFileMetrics options.  Use '' for default options.
    :return:   A tuple with 0) column names and 1) a row iterator.
    """
    sql = embed_parameters(sql, parameters)
    # make automated decision for 'mode'
    if mode == "auto":
        memory_threshold__max_rows = memory_threshold__max_rows or 10000
        memory_threshold__data_size = memory_threshold__data_size or 50000000
        if max_rows and max_rows < memory_threshold__max_rows:
            # low max_rows: we are only looking at the first part of each table
            mode = "pandas"
        else:
            total_size = total_table_stream_sizes(open_table_stream, sql=sql, external_data_indicator=memory_threshold__data_size)
            if total_size < memory_threshold__data_size:
                # total size of all tables is low enough
                mode = "pandas"
            else:
                # large operation, can't run in memory
                mode = "stream"
    if mode == "stream":
        try:
            return streaming_query(sql, open_table_stream, format_options=format_options)
        except Exception as exc:
            raise translate_pandas_dask_exceptions(exc)
    # do in-memory query
    def table_loader(table_spec):
        table_name = table_name_from_table_spec(table_spec)
        table_format = format_options_for_table(format_options, table_spec)
        if max_rows:
            table_format["nrows"] = max_rows
        clean_spec = (table_name,) if isinstance(table_spec, str) else table_spec[:-1] + (table_name,)
        stream = open_table_stream(clean_spec)
        # TODO this metadata might be useful - it is generated in api_sandbox.open_sandbox_file() and other places
        #  where we generate stream classes
        #secret_metadata = (stream.metadata if hasattr(stream, "metadata") else None) or {}
        try:
            return dslibrary.load_dataframe(stream, format_options=table_format, fallback_to_text=True)
        finally:
            stream.close()
        # NOTE: we're not using dask at the moment but if we wanted dslibrary to choose an engine based on file size...
        # , dask=memory_threshold__data_size
    lookup = DbDictLookup(table_loader)
    db = DynamicDict(lookup.tables, lookup.lookup_table)
    return _run_select_s2p(sql, db)


def data_sql_read(data: (str, bytes), sql: str, parameters: list=None, max_rows=None, format_options: dict=None):
    """
    Query raw data as SQL.
    :param data:      Data with columnar data, i.e. CSV, etc.
    :param sql:       SQL to run against the data.
    :param parameters: Parameter values to fill in where '%s' occurs in the SQL.
    :param max_rows:  Limits the amount of raw input data that will be scanned, not the number of output rows.
    :param format_options: See DataFileMetrics.
    """
    def opener(table_spec):
        return io.StringIO(data) if isinstance(data, str) else io.BytesIO(data)
    return file_sql_read(opener, sql=sql, parameters=parameters, max_rows=max_rows, format_options={"": format_options})


def dataframe_sql_read(sql, df):
    """
    Query against a dataframe.
    :param sql:     SQL to run.
    :param df:      A pandas DataFrame, or (cols, rows), or {col: vals, ...}
    :return:  cols, rows
    """
    if isinstance(df, (tuple, list)):
        df = pandas.DataFrame(data=df[1], columns=df[0])
    elif isinstance(df, dict):
        df = pandas.DataFrame(data=df)
    tables = DynamicDict({"table": df}, lambda t: "table")
    return _run_select_s2p(sql, tables)


class DbDictLookup(object):
    def __init__(self, table_loader):
        """
        Generate a table cache and lookup mechanism, given a table loading function.

        :param table_loader:  Does loose matching and retrieves a table, given a table name or
            table specification.
        """
        self.table_loader = table_loader
        self.tables = {}

    def lookup_table(self, table_ref):
        """
        Load a given table, performing loose matching.
        :param table_ref:  Either a str with a table name, or a tuple with a table spec, i.e.
            ("database", "table").
        :returns:       The key under which this table has been loaded, or None if not found.
        """
        table = self.table_loader(table_ref)
        if table is None:
            return
        table_name = table_ref if isinstance(table_ref, str) else table_ref[-1]
        self.tables[table_name] = table
        return table_name


class DynamicDict(object):
    """
    A {} that looks up un-found elements as needed.
    """
    def __init__(self, initial, lookup):
        self.data = {} if initial is None else initial
        self.lookup = lookup

    def get(self, item, default=None):
        found = self.data.get(item)
        if found is not None:
            return found
        new_index = self.lookup(item)
        if new_index is not None:
            return self.data[new_index]
        return default

    def __getitem__(self, item):
        found = self.get(item)
        if found is None:
            raise SqlException_SqlProblem("Table not found: %s" % str(item))
        return found

    def __contains__(self, item):
        return self.get(item) is not None

    def items(self):
        return self.data.items()


def _run_select_s2p(sql, db):
    # references to 'tables' import each table
    ext_tables = lambda rq: (None, "tables[%s]" % repr(rq))
    # capture simple limit - it's a little faster to implement it outside of the main dask loop than with 'head()'
    ext_proc = {"post_limit": None}
    try:
        # compile the SQL
        sql_pandas_code = SqlToPandas(
            sql, local_tables="tables", ext_table_access=ext_tables, external_processing=ext_proc
        ).code
        # execute, filling in tables and dependencies
        vars = {"tables": db, "pandas": pandas, "dask": dask, "numpy": numpy}
        eval(compile(sql_pandas_code, "<sql>", "exec"), vars, vars)
    except Exception as err:
        raise _translate_s2p_exception(err, sql)
    results_df = vars["_"]
    cols = list(results_df)
    # extract rows, and if needed apply post_limit
    row_iter = results_df.itertuples(index=False, name=None)
    if ext_proc["post_limit"]:
        row_iter = slicer(row_iter, slice(0, ext_proc["post_limit"]))
    rows = list(map(DFIterator.clean, row_iter))
    return cols, rows


def _translate_s2p_exception(exc, sql: str):
    """
    Convert errors into code=sql-error where this makes sense.
    """
    if isinstance(exc, SyntaxError):
        return SqlException(f"internal SQL compilation error: {str(exc)}", sql=sql)
    # pandas raises this one
    if exc.__class__.__name__ == "UndefinedVariableError":
        return SqlException_SqlProblem("Undefined column: %s" % str(exc), sql=sql)
    # column not found
    if isinstance(exc, KeyError) and "] are in the [columns]" in str(exc):
        col_info = re.sub(r"None of \[(.+)\] are in the \[columns\]", r'\1', str(exc))
        cols = []
        for col in re.finditer(r'\[([^\]]+)\]', col_info):
            col_list = col.group(1)
            cols += list(map(lambda c: c.strip(" \t'"), col_list.split(",")))
        return SqlException_SqlProblem("Undefined column(s): %s" % (", ".join(cols)), sql=sql)
    if isinstance(exc, KeyError):
        return SqlException_SqlProblem("Unrecognized column: %s" % exc.args[0], sql=sql)
    return exc


def read_sheet_names(input_stream):
    """
    Read sheet names from an XLSX file.
    :returns:  A tuple of column_names, row_values
    """
    book = openpyxl.open(input_stream)
    return ["name"], [[name] for name in book.sheetnames]


def total_table_stream_sizes(open_table_stream: callable, sql: str, external_data_indicator: int=1000000000):
    """
    Determine total size of all referenced tables.
    """
    total = 0
    for table_spec in find_referenced_tables(sql):
        with open_table_stream(table_spec) as stream:
            if isinstance(stream, ChunkedRowIterator):
                return external_data_indicator
            stream.seek(0, 2)
            total += stream.tell()
    return total
